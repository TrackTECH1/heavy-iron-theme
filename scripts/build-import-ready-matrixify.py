#!/usr/bin/env python3
"""Build Matrixify import CSV corrected after Import_Result feedback.

Fixes from Import_Result_2026-06-24_140941.zip:
  - custom.tracktech_data_api must be [json], not [single_line_text_field]
  - Handle = itemid.lower() (canonical)
  - Sprocket images from Supabase manifest (multi-image rows)
  - UTF-8 without BOM on Handle column

Usage:
  python3 scripts/build-import-ready-matrixify.py
  python3 scripts/build-import-ready-matrixify.py --import-result ~/Downloads/Import_Result_2026-06-24_140941.zip
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from lib.catalog_ssot import (  # noqa: E402
    canonical_product_handle,
    catalog_part_family,
    display_tread_pattern,
    is_catalog_itemid,
    load_dotenv_fitment,
    normalize_itemid,
    parse_fitment_csv,
    resolve_track_gallery_urls,
    shopify_product_type,
    shopify_vendor,
    retail_pricing_from_cost,
    product_is_quote_only,
    slugify_make,
    slugify_model,
    supplier_by_itemid,
    normalize_warehouse_availability_json,
    RUBBER_TRACK_WARRANTY_MONTHS,
    supabase_get_all,
)

DEFAULT_IMPORT_ZIP = Path.home() / "Downloads/Import_Result_2026-06-24_140941.zip"
OUTPUT_CATALOG = ROOT / "data/import-ready-catalog-matrixify.csv"
OUTPUT_SPROCKETS = ROOT / "data/sprocket-images-matrixify.csv"
OUTPUT_TRACK_GALLERIES = ROOT / "data/track-gallery-matrixify.csv"
OUTPUT_FITMENT_JSON = ROOT / "data/fitment-json-matrixify.csv"
OUTPUT_FITS_EQUIPMENT_MODELS = ROOT / "data/fits-equipment-models-matrixify.csv"
OUTPUT_UC_FITMENT_JSON = ROOT / "data/uc-fitment-json-matrixify.csv"
OUTPUT_UC_FITS_EQUIPMENT_MODELS = ROOT / "data/uc-fits-equipment-models-matrixify.csv"
OUTPUT_GAP_REPORT = ROOT / "data/fitment-gap-report.csv"
MANIFEST = ROOT / "data/uc-images-manifest.json"
LEGACY_MANIFEST = ROOT / "data/sprocket-images-manifest.json"

MATRIXIFY_COLUMNS = [
    "Handle",
    "Command",
    "Title",
    "Body HTML",
    "Vendor",
    "Type",
    "Tags",
    "Status",
    "Published",
    "Top Row",
    "Image Src",
    "Image Position",
    "Variant SKU",
    "Variant Price",
    "Variant Compare At Price",
    "Variant Cost",
    "Variant Weight",
    "Variant Weight Unit",
    "Variant Inventory Policy",
    "Variant Inventory Tracker",
    "Metafield: custom.tracktech_data_api [json]",
    "Metafield: custom.track_size [single_line_text_field]",
    "Metafield: custom.tread_pattern [single_line_text_field]",
    "Metafield: custom.mpn [single_line_text_field]",
    "Metafield: custom.part_type [single_line_text_field]",
    "Metafield: custom.width_mm [number_integer]",
    "Metafield: custom.pitch_mm [number_decimal]",
    "Metafield: custom.links [number_integer]",
    "Metafield: custom.warranty_months [number_integer]",
    "Metafield: custom.fitment_json [json]",
    "Variant Metafield: custom.warehouse_availability [json]",
]

FITMENT_JSON_COLUMNS = [
    "Handle",
    "Command",
    "Top Row",
    "Variant SKU",
    "Metafield: custom.fitment_json [json]",
]

FITS_EQUIPMENT_MODELS_COLUMNS = [
    "Handle",
    "Command",
    "Top Row",
    "Variant SKU",
    "Metafield: custom.fits_equipment_models [list.metaobject_reference]",
]

MODEL_METAOBJECT_TYPE = "model"

MACHINE_TYPE_LABELS = {
    "compact_track_loader": "Compact Track Loader",
    "mini_track_loader": "Mini Track Loader",
    "mini_excavator": "Mini Excavator",
    "compact_utility_loader": "Compact Utility Loader",
}


def shopify_tags(itemid: str, family: str, track_size: str | None, *, quote_only: bool = False) -> str:
    tags = ["itemid-ssot", f"itemid:{itemid}"]
    if quote_only:
        tags.append("request-quote")
    if family == "track":
        tags.extend(["rubber-tracks", "nav:rubber-tracks"])
        if track_size:
            tags.append(f"size:{track_size}")
    else:
        tags.append("undercarriage-parts")
        if family in {"sprocket", "idler", "roller"}:
            tags.append(f"uc:{family}")
    return ", ".join(tags)


def build_tracktech_json(product: dict, *, family: str = "") -> str:
    itemid = product["itemid"]
    payload: dict = {"itemid": itemid}
    if product.get("title"):
        payload["product_name"] = product["title"]
    if product.get("price") is not None:
        payload["price"] = float(product["price"])
    if family == "track":
        payload["warranty_months"] = RUBBER_TRACK_WARRANTY_MONTHS
    if product.get("track_size"):
        payload["spec_track_size"] = product["track_size"]
    tread = display_tread_pattern(product.get("tread_pattern")) or product.get("tread_pattern")
    if tread:
        payload["spec_track_pattern"] = tread
    if product.get("width_mm") is not None:
        payload["spec_width_mm"] = str(int(float(product["width_mm"])))
    if product.get("pitch_mm") is not None:
        payload["spec_pitch"] = str(product["pitch_mm"])
    if product.get("links") is not None:
        payload["spec_links"] = str(int(product["links"]))
    mpn = product.get("mpn") or product.get("oem_part_number") or itemid
    payload["spec_mpn"] = mpn
    return json.dumps(payload, separators=(",", ":"))


def build_rubber_track_body_html(product: dict) -> str:
    itemid = product["itemid"]
    track_size = product.get("track_size") or ""
    tread = display_tread_pattern(product.get("tread_pattern")) or product.get("tread_pattern") or ""
    width = product.get("width_mm")
    pitch = product.get("pitch_mm")
    links = product.get("links")
    weight = product.get("weight_lbs")

    parts: list[str] = []
    if track_size:
        parts.append(f"sized <strong>{track_size}</strong>")
    dims: list[str] = []
    if width is not None:
        dims.append(f"{int(float(width))}&nbsp;mm width")
    if pitch is not None:
        dims.append(f"{pitch}&nbsp;mm pitch")
    if links is not None:
        dims.append(f"{int(float(links))} links")
    if dims:
        parts.append(f"({' &times; '.join(dims)})")

    lead = " ".join(parts).strip()
    if tread and lead:
        opener = f"<p>Heavy-duty {tread} rubber track {lead}."
    elif lead:
        opener = f"<p>OEM-spec rubber track {lead}."
    else:
        opener = "<p>OEM-spec rubber track built for demanding jobsite conditions."

    body = (
        f"{opener} Continuous steel cord construction, forged steel links, and an "
        "abrasion-resistant rubber compound deliver long service life with guaranteed fitment.</p>"
        f"<p>{RUBBER_TRACK_WARRANTY_MONTHS}-month manufacturer warranty. Free LTL freight to the contiguous US. "
        "Same-day shipping on orders placed before 12&nbsp;PM&nbsp;CT.</p>"
    )
    bullets: list[str] = []
    if tread:
        bullets.append(f"<li><strong>Pattern:</strong> {tread}</li>")
    if weight is not None:
        w = float(weight)
        w_s = str(int(w)) if w == int(w) else str(w)
        bullets.append(f"<li><strong>Weight:</strong> {w_s} lbs per track</li>")
    bullets.append(f"<li><strong>SKU:</strong> {itemid}</li>")
    return body + f"<ul>{''.join(bullets)}</ul>"


def humanize_machine_type(code: str | None) -> str:
    raw = (code or "").strip()
    if not raw:
        return ""
    if raw in MACHINE_TYPE_LABELS:
        return MACHINE_TYPE_LABELS[raw]
    return raw.replace("_", " ").title()


def load_store_skus(export_path: Path) -> set[str]:
    """Variant SKUs from a Matrixify Products export (live store catalog)."""
    suffix = export_path.suffix.lower()
    skus: set[str] = set()
    if suffix == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(export_path, read_only=True, data_only=True)
        ws = wb["Products"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {name: i for i, name in enumerate(headers) if name}
        sku_i = idx.get("Variant SKU")
        if sku_i is None:
            wb.close()
            raise ValueError(f"No Variant SKU column in {export_path}")
        for row in ws.iter_rows(min_row=2, values_only=True):
            sku = normalize_itemid(row[sku_i] if sku_i < len(row) else None)
            if sku:
                skus.add(sku)
        wb.close()
        return skus

    with export_path.open(newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            sku = normalize_itemid(row.get("Variant SKU") or row.get("Handle"))
            if sku:
                skus.add(sku)
    return skus


def load_fitments_by_itemid(*, fit_type: str = "track") -> dict[str, list[dict]]:
    """Supabase fitments grouped by catalog itemid."""
    rows = supabase_get_all(
        "fitment?select=source,product:product_id(itemid),"
        f"model:model_id(make,model,machine_type_code,model_key)"
        f"&fit_type=eq.{fit_type}"
    )
    by_itemid: dict[str, list[dict]] = {}
    seen: dict[str, set[str]] = {}
    for row in rows:
        product = row.get("product") or {}
        model = row.get("model") or {}
        itemid = normalize_itemid(product.get("itemid"))
        model_key = (model.get("model_key") or "").strip()
        if not itemid or not model_key:
            continue
        dedupe_key = model_key.lower()
        if dedupe_key in seen.setdefault(itemid, set()):
            continue
        seen[itemid].add(dedupe_key)
        by_itemid.setdefault(itemid, []).append(
            {
                "make": (model.get("make") or "").strip(),
                "model": (model.get("model") or "").strip(),
                "machine_type": humanize_machine_type(model.get("machine_type_code")),
                "machine_type_code": (model.get("machine_type_code") or "").strip(),
                "model_key": model_key,
                "source": (row.get("source") or "").strip(),
            }
        )
    for fits in by_itemid.values():
        fits.sort(key=lambda f: (f.get("make") or "", f.get("model") or ""))
    return by_itemid


def build_fitment_display_text(fits: list[dict], track_size: str, machine_type: str) -> str:
    if not fits:
        return ""
    makes: list[str] = []
    seen: set[str] = set()
    for fit in fits:
        make = (fit.get("make") or "").strip()
        key = make.lower()
        if make and key not in seen:
            seen.add(key)
            makes.append(make)
    if not makes:
        return ""
    if len(makes) == 1:
        make_part = makes[0]
    elif len(makes) <= 4:
        make_part = ", ".join(makes[:-1]) + f", and {makes[-1]}"
    else:
        make_part = ", ".join(makes[:3]) + f", and {len(makes) - 3} other brands"
    type_label = (machine_type or "machines").lower()
    size_bit = f" using {track_size} rubber tracks" if track_size else ""
    return f"Fits select {make_part} and other {type_label}{size_bit}."


def build_fitment_json_from_supplier(product: dict, supplier_row: dict | None) -> str:
    raw = str((supplier_row or {}).get("fitment_models") or "").strip()
    if not raw:
        return ""
    track_size = (product.get("track_size") or "").strip()
    fits: list[dict] = []
    for make, model_name in parse_fitment_csv(raw):
        fits.append(
            {
                "make": make.strip(),
                "model": model_name.strip(),
                "machine_type": "",
                "machine_type_code": "",
                "model_key": f"{slugify_make(make)}-{slugify_model(model_name)}",
                "source": "supplier_fitment",
            }
        )
    if not fits:
        return ""
    return build_fitment_json(product, fits, verified=False)


def build_fitment_json(product: dict, fits: list[dict], *, verified: bool | None = None) -> str:
    if not fits:
        return ""
    track_size = (product.get("track_size") or "").strip()
    machine_types = [f.get("machine_type") for f in fits if f.get("machine_type")]
    dominant_type = ""
    if machine_types:
        dominant_type = max(set(machine_types), key=machine_types.count)
    verified_sources = {
        "tracktech_export",
        "explicit_supplier_fitment",
        "supplier_row_model",
        "supplier_fitment",
    }
    if verified is None:
        verified = any((f.get("source") or "") in verified_sources for f in fits) or len(fits) > 0
    payload = {
        "verified": verified,
        "machine_type": dominant_type,
        "track_size": track_size,
        "fits": [
            {
                "make": f.get("make") or "",
                "model": f.get("model") or "",
                "machine_type": f.get("machine_type") or "",
                "track_size": track_size,
                "model_key": f.get("model_key") or "",
            }
            for f in fits
        ],
        "display_text": build_fitment_display_text(fits, track_size, dominant_type),
    }
    return json.dumps(payload, separators=(",", ":"))


def load_model_handles_by_itemid(*, fit_type: str = "track") -> dict[str, list[str]]:
    """Published model metaobject handles per SKU (requires shopify_metaobject_gid)."""
    rows = supabase_get_all(
        "fitment?select=product:product_id(itemid),"
        f"model:model_id(model_key,model_handle,shopify_metaobject_gid)"
        f"&fit_type=eq.{fit_type}"
    )
    by_itemid: dict[str, list[str]] = {}
    seen: dict[str, set[str]] = {}
    for row in rows:
        product = row.get("product") or {}
        model = row.get("model") or {}
        itemid = normalize_itemid(product.get("itemid"))
        if not itemid or not (model.get("shopify_metaobject_gid") or "").strip():
            continue
        handle = (model.get("model_handle") or model.get("model_key") or "").strip()
        if not handle:
            continue
        key = handle.lower()
        if key in seen.setdefault(itemid, set()):
            continue
        seen[itemid].add(key)
        by_itemid.setdefault(itemid, []).append(handle)
    for handles in by_itemid.values():
        handles.sort(key=str.lower)
    return by_itemid


def format_metaobject_reference_list(handles: list[str], definition: str = MODEL_METAOBJECT_TYPE) -> str:
    return ", ".join(f"{definition}.{handle}" for handle in handles)


def load_products() -> list[dict]:
    rows = supabase_get_all(
        "product?select=id,sku,itemid,handle,title,part_type,type,track_size,tread_pattern,"
        "price,cost,weight_lbs,mpn,oem_part_number,part_number,width_mm,pitch_mm,links,image_url"
        "&source=eq.itemid_ssot&order=itemid.asc"
    )
    out: list[dict] = []
    seen: set[str] = set()
    for r in rows:
        itemid = normalize_itemid(r.get("itemid") or r.get("sku"))
        if not itemid or not is_catalog_itemid(itemid) or itemid in seen:
            continue
        seen.add(itemid)
        family = catalog_part_family(itemid, r.get("part_type") or r.get("type"))
        if family not in {"track", "sprocket", "idler", "roller"}:
            continue
        out.append(
            {
                **r,
                "itemid": itemid,
                "family": family,
                "canonical_handle": canonical_product_handle(itemid),
            }
        )
    return out


def load_uc_images() -> dict[str, list[str]]:
    path = MANIFEST if MANIFEST.is_file() else LEGACY_MANIFEST
    if not path.is_file():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    out: dict[str, list[str]] = {}
    for entry in data:
        itemid = normalize_itemid(entry.get("itemid"))
        urls = [img["public_url"] for img in entry.get("images") or [] if img.get("public_url")]
        if itemid and urls:
            out[itemid] = urls
    return out


def load_sprocket_images() -> dict[str, list[str]]:
    return load_uc_images()


def load_import_result_skus(zip_path: Path) -> set[str]:
    """SKUs that already imported OK (for reference only)."""
    if not zip_path.is_file():
        return set()
    with zipfile.ZipFile(zip_path) as zf:
        name = next(n for n in zf.namelist() if n.endswith(".csv"))
        with zf.open(name) as f:
            text = f.read().decode("utf-8-sig")
    rows = list(csv.DictReader(text.splitlines()))
    return {
        normalize_itemid(r.get("Variant SKU") or r.get("Handle"))
        for r in rows
        if r.get("Import Result") == "OK" and (r.get("Variant SKU") or r.get("Handle") or "").strip()
    }


def matrixify_product_row(
    product: dict,
    *,
    image_src: str = "",
    image_position: int = 0,
    is_top_row: bool = False,
    fitment_json: str = "",
) -> dict[str, str]:
    itemid = product["itemid"]
    family = product["family"]
    row = {col: "" for col in MATRIXIFY_COLUMNS}
    row["Handle"] = product["canonical_handle"]
    if is_top_row:
        row["Command"] = "MERGE"
        row["Title"] = product.get("title") or itemid
        row["Vendor"] = shopify_vendor(itemid, family)
        row["Type"] = shopify_product_type(family)
        if family == "track":
            row["Body HTML"] = build_rubber_track_body_html(product)
            row["Metafield: custom.warranty_months [number_integer]"] = str(RUBBER_TRACK_WARRANTY_MONTHS)
        quote_only = product_is_quote_only(product.get("cost"))
        row["Tags"] = shopify_tags(itemid, family, product.get("track_size"), quote_only=quote_only)
        row["Status"] = "Active"
        row["Published"] = "TRUE"
        row["Top Row"] = "TRUE"
        row["Variant SKU"] = itemid
        pricing = retail_pricing_from_cost(product.get("cost"))
        if pricing:
            row["Variant Cost"], row["Variant Price"], row["Variant Compare At Price"] = pricing
            row["Variant Inventory Policy"] = "continue"
        else:
            row["Variant Inventory Policy"] = "deny"
        if product.get("weight_lbs") is not None:
            w = float(product["weight_lbs"])
            row["Variant Weight"] = str(int(w)) if w == int(w) else str(w)
            row["Variant Weight Unit"] = "lb"
        row["Metafield: custom.tracktech_data_api [json]"] = build_tracktech_json(product, family=family)
        if product.get("track_size"):
            row["Metafield: custom.track_size [single_line_text_field]"] = product["track_size"]
        tread = display_tread_pattern(product.get("tread_pattern")) or product.get("tread_pattern") or ""
        if tread:
            row["Metafield: custom.tread_pattern [single_line_text_field]"] = tread
        mpn = product.get("mpn") or product.get("oem_part_number") or itemid
        row["Metafield: custom.mpn [single_line_text_field]"] = mpn
        if product.get("part_type"):
            row["Metafield: custom.part_type [single_line_text_field]"] = product["part_type"]
        if product.get("width_mm") is not None:
            row["Metafield: custom.width_mm [number_integer]"] = str(int(float(product["width_mm"])))
        if product.get("pitch_mm") is not None:
            row["Metafield: custom.pitch_mm [number_decimal]"] = str(product["pitch_mm"])
        if product.get("links") is not None:
            row["Metafield: custom.links [number_integer]"] = str(int(product["links"]))
        warehouse_json = normalize_warehouse_availability_json(product.get("warehouse_availability"))
        if warehouse_json:
            row["Variant Metafield: custom.warehouse_availability [json]"] = warehouse_json
        if family == "track" and fitment_json:
            row["Metafield: custom.fitment_json [json]"] = fitment_json
    if image_src:
        row["Image Src"] = image_src
        row["Image Position"] = str(image_position or 1)
    return row


def product_image_urls(product: dict, uc_images: dict[str, list[str]]) -> list[str]:
    itemid = product["itemid"]
    family = product["family"]
    if family in {"sprocket", "idler", "roller"} and itemid in uc_images:
        return uc_images[itemid]
    if family == "track":
        tread = display_tread_pattern(product.get("tread_pattern")) or product.get("tread_pattern")
        urls = resolve_track_gallery_urls(itemid, product.get("track_size"), tread)
        if urls:
            return urls
    img = (product.get("image_url") or "").strip()
    return [img] if img else []


def rows_for_product(product: dict, urls: list[str], *, fitment_json: str = "") -> list[dict[str, str]]:
    images = urls or [""]
    rows: list[dict[str, str]] = []
    for idx, url in enumerate(images):
        rows.append(
            matrixify_product_row(
                product,
                image_src=url,
                image_position=idx + 1 if url else 0,
                is_top_row=idx == 0,
                fitment_json=fitment_json,
            )
        )
    return rows


def fitment_json_row(product: dict, fitment_json: str) -> dict[str, str]:
    row = {col: "" for col in FITMENT_JSON_COLUMNS}
    row["Handle"] = product["canonical_handle"]
    row["Command"] = "MERGE"
    row["Top Row"] = "TRUE"
    row["Variant SKU"] = product["itemid"]
    row["Metafield: custom.fitment_json [json]"] = fitment_json
    return row


def fits_equipment_models_row(product: dict, refs: str) -> dict[str, str]:
    row = {col: "" for col in FITS_EQUIPMENT_MODELS_COLUMNS}
    row["Handle"] = product["canonical_handle"]
    row["Command"] = "MERGE"
    row["Top Row"] = "TRUE"
    row["Variant SKU"] = product["itemid"]
    row["Metafield: custom.fits_equipment_models [list.metaobject_reference]"] = refs
    return row


def _gap_reason(
    fits: list[dict],
    fitment_json: str,
    refs: str,
    supplier_row: dict | None,
) -> str:
    if fitment_json and refs:
        return ""
    reasons: list[str] = []
    if not fitment_json:
        if not fits and not str((supplier_row or {}).get("fitment_models") or "").strip():
            reasons.append("no_supplier_fitment")
        elif not fits:
            reasons.append("supplier_unparsed")
        else:
            reasons.append("no_fitment_json")
    if not refs:
        if not fits:
            reasons.append("no_fitment_rows")
        else:
            reasons.append("models_missing_shopify_gid")
    return ";".join(reasons)


def write_csv(path: Path, rows: list[dict[str, str]], *, columns: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = columns or MATRIXIFY_COLUMNS
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build import-ready Matrixify CSV")
    parser.add_argument("--import-result", type=Path, default=DEFAULT_IMPORT_ZIP)
    parser.add_argument("--catalog-out", type=Path, default=OUTPUT_CATALOG)
    parser.add_argument("--sprocket-out", type=Path, default=OUTPUT_SPROCKETS)
    parser.add_argument("--track-gallery-out", type=Path, default=OUTPUT_TRACK_GALLERIES)
    parser.add_argument("--fitment-json-out", type=Path, default=OUTPUT_FITMENT_JSON)
    parser.add_argument("--fits-equipment-models-out", type=Path, default=OUTPUT_FITS_EQUIPMENT_MODELS)
    parser.add_argument("--uc-fitment-json-out", type=Path, default=OUTPUT_UC_FITMENT_JSON)
    parser.add_argument("--uc-fits-equipment-models-out", type=Path, default=OUTPUT_UC_FITS_EQUIPMENT_MODELS)
    parser.add_argument("--gap-report-out", type=Path, default=OUTPUT_GAP_REPORT)
    parser.add_argument(
        "--store-export",
        type=Path,
        default=None,
        help="Limit fitment Matrixify rows to Variant SKUs present in a live store export",
    )
    args = parser.parse_args()

    load_dotenv_fitment()
    products = load_products()
    store_skus: set[str] | None = None
    if args.store_export:
        if not args.store_export.is_file():
            print(f"Store export not found: {args.store_export}", file=sys.stderr)
            return 1
        store_skus = load_store_skus(args.store_export)
        products = [p for p in products if p["itemid"] in store_skus]
        print(f"Store export filter: {len(products)} catalog SKUs in {args.store_export.name}")

    track_fitments_by_itemid = load_fitments_by_itemid(fit_type="track")
    uc_fitments_by_itemid = load_fitments_by_itemid(fit_type="uc_part")
    track_model_handles_by_itemid = load_model_handles_by_itemid(fit_type="track")
    uc_model_handles_by_itemid = load_model_handles_by_itemid(fit_type="uc_part")
    supplier = supplier_by_itemid()
    for p in products:
        sup = supplier.get(p["itemid"]) or {}
        if sup.get("warehouse_availability"):
            p["warehouse_availability"] = sup["warehouse_availability"]
    uc_images = load_uc_images()
    ok_skus = load_import_result_skus(args.import_result)

    catalog_rows: list[dict[str, str]] = []
    sprocket_only_rows: list[dict[str, str]] = []
    track_gallery_rows: list[dict[str, str]] = []
    fitment_json_rows: list[dict[str, str]] = []
    fits_equipment_models_rows: list[dict[str, str]] = []
    uc_fitment_json_rows: list[dict[str, str]] = []
    uc_fits_equipment_models_rows: list[dict[str, str]] = []
    gap_report_rows: list[dict[str, str]] = []
    model_ref_counts: list[int] = []
    uc_model_ref_counts: list[int] = []

    for p in sorted(products, key=lambda x: x["itemid"]):
        itemid = p["itemid"]
        urls = product_image_urls(p, uc_images)
        family = p["family"]
        fitment_json = ""
        fits = []
        model_handles: list[str] = []

        if family == "track":
            fits = track_fitments_by_itemid.get(itemid, [])
            fitment_json = build_fitment_json(p, fits) if fits else build_fitment_json_from_supplier(p, supplier.get(itemid))
            if fitment_json:
                fitment_json_rows.append(fitment_json_row(p, fitment_json))
            model_handles = track_model_handles_by_itemid.get(itemid, [])
            refs = format_metaobject_reference_list(model_handles)
            if refs:
                fits_equipment_models_rows.append(fits_equipment_models_row(p, refs))
                model_ref_counts.append(len(model_handles))
            gap_report_rows.append(
                {
                    "itemid": itemid,
                    "family": family,
                    "fitment_rows": str(len(fits)),
                    "fitment_json": "yes" if fitment_json else "no",
                    "fits_equipment_models": "yes" if refs else "no",
                    "model_ref_count": str(len(model_handles)),
                    "gap_reason": _gap_reason(fits, fitment_json, refs, supplier.get(itemid)),
                }
            )
        elif family in {"sprocket", "idler", "roller"}:
            fits = uc_fitments_by_itemid.get(itemid, [])
            fitment_json = build_fitment_json(p, fits) if fits else build_fitment_json_from_supplier(p, supplier.get(itemid))
            if fitment_json:
                uc_fitment_json_rows.append(fitment_json_row(p, fitment_json))
            model_handles = uc_model_handles_by_itemid.get(itemid, [])
            refs = format_metaobject_reference_list(model_handles)
            if refs:
                uc_fits_equipment_models_rows.append(fits_equipment_models_row(p, refs))
                uc_model_ref_counts.append(len(model_handles))
            gap_report_rows.append(
                {
                    "itemid": itemid,
                    "family": family,
                    "fitment_rows": str(len(fits)),
                    "fitment_json": "yes" if fitment_json else "no",
                    "fits_equipment_models": "yes" if refs else "no",
                    "model_ref_count": str(len(model_handles)),
                    "gap_reason": _gap_reason(fits, fitment_json, refs, supplier.get(itemid)),
                }
            )

        catalog_rows.extend(rows_for_product(p, urls, fitment_json=fitment_json if family == "track" else ""))

        if p["family"] in {"sprocket", "idler", "roller"} and urls:
            sprocket_only_rows.extend(rows_for_product(p, urls))
        elif p["family"] == "track" and len(urls) > 1:
            track_gallery_rows.extend(rows_for_product(p, urls))

    write_csv(args.catalog_out, catalog_rows)
    write_csv(args.sprocket_out, sprocket_only_rows)
    write_csv(args.track_gallery_out, track_gallery_rows)
    write_csv(args.fitment_json_out, fitment_json_rows, columns=FITMENT_JSON_COLUMNS)
    write_csv(args.fits_equipment_models_out, fits_equipment_models_rows, columns=FITS_EQUIPMENT_MODELS_COLUMNS)
    write_csv(args.uc_fitment_json_out, uc_fitment_json_rows, columns=FITMENT_JSON_COLUMNS)
    write_csv(args.uc_fits_equipment_models_out, uc_fits_equipment_models_rows, columns=FITS_EQUIPMENT_MODELS_COLUMNS)
    write_csv(
        args.gap_report_out,
        gap_report_rows,
        columns=[
            "itemid",
            "family",
            "fitment_rows",
            "fitment_json",
            "fits_equipment_models",
            "model_ref_count",
            "gap_reason",
        ],
    )

    from collections import Counter

    families = Counter(p["family"] for p in products)
    track_count = sum(1 for p in products if p["family"] == "track")
    print(f"Wrote {args.catalog_out} ({len(catalog_rows)} rows, {len(products)} products)")
    print(f"Wrote {args.sprocket_out} ({len(sprocket_only_rows)} rows)")
    track_products = sum(1 for p in products if p["family"] == "track")
    track_multi = sum(1 for p in products if p["family"] == "track" and len(product_image_urls(p, uc_images)) > 1)
    print(f"Wrote {args.track_gallery_out} ({len(track_gallery_rows)} rows, {track_multi}/{track_products} tracks with galleries)")
    print(f"Wrote {args.fitment_json_out} ({len(fitment_json_rows)} rows, {len(fitment_json_rows)}/{track_count} tracks with fitment_json)")
    if model_ref_counts:
        print(
            f"Wrote {args.fits_equipment_models_out} ({len(fits_equipment_models_rows)} rows, "
            f"max {max(model_ref_counts)} model refs/product, "
            f"{sum(model_ref_counts)} total refs)"
        )
    else:
        print(f"Wrote {args.fits_equipment_models_out} (0 rows)")
    uc_count = sum(1 for p in products if p["family"] in {"sprocket", "idler", "roller"})
    print(
        f"Wrote {args.uc_fitment_json_out} ({len(uc_fitment_json_rows)} rows, "
        f"{len(uc_fitment_json_rows)}/{uc_count} UC with fitment_json)"
    )
    if uc_model_ref_counts:
        print(
            f"Wrote {args.uc_fits_equipment_models_out} ({len(uc_fits_equipment_models_rows)} rows, "
            f"{sum(uc_model_ref_counts)} total refs)"
        )
    else:
        print(f"Wrote {args.uc_fits_equipment_models_out} (0 rows)")
    gaps = [row for row in gap_report_rows if row["gap_reason"]]
    print(f"Wrote {args.gap_report_out} ({len(gap_report_rows)} rows, {len(gaps)} with remaining gaps)")
    print(f"  families: {dict(families)}")
    print(f"  prior import OK SKUs: {len(ok_skus)}")
    print("Matrixify: MERGE, identify by Variant SKU, enable Products/Variants/Images/Metafields")
    print("Fitment JSON pass: MERGE by Variant SKU, enable Products + Metafields only")
    print("Fits Models pass: MERGE by Variant SKU, enable Products + Metafields only (model.handle refs)")
    print("Fix: tracktech_data_api uses [json] column with valid JSON object")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
