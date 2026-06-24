#!/usr/bin/env python3
"""Matrixify DELETE rows for excluded catalog itemids in a live store export.

Usage:
  python3 scripts/build-excluded-delete-matrixify.py --export ~/Downloads/Export_*.xlsx
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from lib.catalog_ssot import is_excluded_catalog_itemid, load_dotenv_fitment, normalize_itemid  # noqa: E402

DEFAULT_EXPORT = Path.home() / "Downloads/Export_2026-06-24_160525.xlsx"
OUTPUT = ROOT / "data/excluded-solid-tire-delete-matrixify.csv"


def load_export_rows(export_path: Path) -> list[dict[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(export_path, read_only=True, data_only=True)
    ws = wb["Products"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(headers) if name}
    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[idx.get("Top Row", -1)] or "").upper() != "TRUE":
            continue

        def col(name: str) -> str:
            i = idx.get(name)
            if i is None or i >= len(row):
                return ""
            return str(row[i] or "").strip()

        sku = normalize_itemid(col("Variant SKU"))
        if not sku or not is_excluded_catalog_itemid(sku):
            continue
        rows.append(
            {
                "ID": col("ID"),
                "Handle": col("Handle").lower(),
                "Command": "DELETE",
                "Variant SKU": sku,
            }
        )
    wb.close()
    return rows


def main() -> int:
    load_dotenv_fitment()
    parser = argparse.ArgumentParser(description="Build Matrixify DELETE for excluded catalog SKUs")
    parser.add_argument("--export", type=Path, default=DEFAULT_EXPORT)
    parser.add_argument("--out", type=Path, default=OUTPUT)
    args = parser.parse_args()
    if not args.export.is_file():
        print(f"Export not found: {args.export}", file=sys.stderr)
        return 1

    rows = load_export_rows(args.export)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["ID", "Handle", "Command", "Variant SKU"])
        writer.writeheader()
        writer.writerows(rows)

    print(f"Wrote {args.out} ({len(rows)} DELETE rows)")
    print("Matrixify: Delete mode, identify by ID or Handle, Products enabled only")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
